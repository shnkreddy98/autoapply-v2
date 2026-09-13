"""Onboarding (`autoapply profile init`): extraction, LLM parse plumbing,
interactive review defaults, and the DB/storage save. LLM is faked throughout."""

import json
from datetime import date

import pytest

from cli import onboard
from cli.onboard import (
    ParsedExperience,
    ParsedResume,
    interactive_review,
    parse_resume_llm,
    save,
    years_of_experience,
)

# ---------------------------------------------------------------------------
# extraction


def test_extract_text_docx(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Data Engineer at Acme")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "SQL"
    path = tmp_path / "r.docx"
    doc.save(str(path))

    text = onboard.extract_text(str(path))
    assert "Jane Doe" in text
    assert "Data Engineer at Acme" in text
    assert "Python\tSQL" in text


def test_docx_to_pdf_conversion(tmp_path):
    import shutil

    if not shutil.which("libreoffice"):
        pytest.skip("LibreOffice not installed")
    import docx

    doc = docx.Document()
    doc.add_paragraph("Jane Doe — Data Engineer")
    src = tmp_path / "r.docx"
    doc.save(str(src))

    pdf = onboard.docx_to_pdf(str(src))
    assert pdf.endswith("r.pdf")
    assert open(pdf, "rb").read(4) == b"%PDF"


def test_extract_text_rejects_unknown_format(tmp_path):
    p = tmp_path / "r.txt"
    p.write_text("hi")
    with pytest.raises(SystemExit, match="unsupported"):
        onboard.extract_text(str(p))


# ---------------------------------------------------------------------------
# LLM parse (fake client)

_GOOD = {
    "contact": {
        "full_name": "Jane Doe",
        "email": "jane@x.dev",
        "location": "Austin, TX",
    },
    "summary": "Data engineer.",
    "work_experiences": [
        {
            "job_title": "DE",
            "company_name": "Acme",
            "from_date": "2020-01-01",
            "is_current": True,
            "bullets": ["built pipelines"],
        }
    ],
    "education": [{"degree": "MS", "school": "UT", "gpa": "3.9/4.0"}],
    "skill_groups": [{"category": "Languages", "skills": ["Python", "SQL"]}],
}


class FakeClient:
    def __init__(self, replies):
        self._replies = list(replies)
        self.calls = 0

        outer = self

        class _Completions:
            def create(self, **kw):
                outer.calls += 1
                reply = outer._replies.pop(0)

                class Msg:
                    content = reply

                class Choice:
                    message = Msg()

                class Resp:
                    choices = [Choice()]

                return Resp()

        class _Chat:
            completions = _Completions()

        self.chat = _Chat()


def test_parse_resume_llm_happy_path():
    parsed = parse_resume_llm("resume text", client=FakeClient([json.dumps(_GOOD)]))
    assert parsed.contact.email == "jane@x.dev"
    assert parsed.work_experiences[0].is_current is True
    assert parsed.work_experiences[0].to_date is None
    assert parsed.skill_groups[0].skills == ["Python", "SQL"]


def test_parse_resume_llm_strips_fences_and_retries():
    fenced = f"```json\n{json.dumps(_GOOD)}\n```"
    parsed = parse_resume_llm("txt", client=FakeClient([fenced]))
    assert parsed.contact.full_name == "Jane Doe"

    client = FakeClient(["not json at all", json.dumps(_GOOD)])
    parsed = parse_resume_llm("txt", client=client)
    assert client.calls == 2
    assert parsed.contact.email == "jane@x.dev"


def test_parse_resume_llm_gives_up_after_retry():
    with pytest.raises(SystemExit, match="parse failed"):
        parse_resume_llm("txt", client=FakeClient(["nope", "still nope"]))


# ---------------------------------------------------------------------------
# review


def _parsed() -> ParsedResume:
    return ParsedResume.model_validate(_GOOD)


def test_years_of_experience():
    assert years_of_experience([], today=date(2026, 1, 1)) is None
    work = [
        ParsedExperience(from_date=date(2020, 1, 1)),
        ParsedExperience(from_date=date(2023, 1, 1)),
    ]
    assert years_of_experience(work, today=date(2026, 1, 1)) == 6


def test_review_yes_accepts_all_defaults():
    # the password is the one prompt --yes does NOT skip (it has no default)
    data = interactive_review(_parsed(), yes=True, ask_secret=lambda _: "hunter2")
    assert data.resume.contact.email == "jane@x.dev"
    assert data.login == {"login_email": "jane@x.dev", "login_password": "hunter2"}
    p = data.profile
    assert p["city"] == "Austin" and p["state"] == "TX"  # split from location
    assert p["work_eligible_us"] is True
    assert p["visa_sponsorship"] is False
    assert p["gender"] == "Prefer not to say"
    assert p["cert_accuracy"] is True and p["cert_data_processing"] is True
    assert p["electronic_signature"] == "Jane Doe"
    assert p["signature_date"] == date.today()
    assert p["years_of_experience"] >= 6
    assert (p["first_name"], p["middle_name"], p["last_name"]) == ("Jane", None, "Doe")
    assert p["phone_country_code"] == "+1" and p["country"] == "United States"
    assert p["over_18"] is True
    assert p["hispanic_latino"] == "Prefer not to say"
    assert p["notice_period"] == "2 weeks" and p["how_did_you_hear"] == "LinkedIn"


def test_split_name_defaults():
    """Only a default for the review prompt — compound surnames split wrong, which
    is why the confirmed parts are stored instead of derived at fill time."""
    assert onboard._split_name("Jane Quinn Doe") == ("Jane", "Quinn", "Doe")
    assert onboard._split_name("Cher") == ("Cher", None, None)
    assert onboard._split_name(None) == (None, None, None)


def test_review_password_skippable():
    data = interactive_review(_parsed(), yes=True, ask_secret=lambda _: "  ")
    assert data.login["login_password"] is None


def test_review_prompts_override_and_clear():
    answers = iter(
        [
            "Janet Doe",  # full name overridden
            "",  # email kept
            "-",  # phone cleared
            "",
            "",
            "",  # location, linkedin, github kept
            "",  # first name kept (split from the overridden full name)
            "",  # middle name blank
            "Doe-Smith",  # last name overridden
            "",
            "",
            "",
            "",  # preferred, pronouns, phone code, portfolio kept
            "portal@x.dev",  # portal login email overridden
            "",
            "",
            "",
            "",
            "",
            "",  # address block (street, city, state, zip, country, dob) kept/blank
            "",
            "y",
            "",  # eligible(kept) sponsorship(yes) clearance(kept)
            "",  # over 18 kept
            "10",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",  # yoe=10, rest defaults (through notice period, how did you hear)
            "",  # ever terminated (kept: no)
            "",
            "",
            "",
            "",
            "",  # EEO
            "",  # certs agree
            "",  # signature
        ]
    )
    data = interactive_review(
        _parsed(), ask=lambda _: next(answers), ask_secret=lambda _: "pw1"
    )
    assert data.resume.contact.full_name == "Janet Doe"
    assert data.resume.contact.email == "jane@x.dev"
    assert data.resume.contact.phone is None
    assert data.profile["visa_sponsorship"] is True
    assert data.profile["years_of_experience"] == 10
    assert data.profile["first_name"] == "Janet"
    assert data.profile["last_name"] == "Doe-Smith"
    assert data.login == {"login_email": "portal@x.dev", "login_password": "pw1"}


def test_review_yes_keeps_existing_profile_values():
    # re-running init with --yes must not wipe previously entered answers
    existing = {
        "street_address": "100 Example St",
        "zip_code": "95126",
        "city": "San Jose",  # beats the resume-location guess (Austin)
        "visa_sponsorship": True,
        "years_of_experience": 9,
        "gender": "Woman",
        "desired_salary": "150000",
    }
    data = interactive_review(
        _parsed(), yes=True, ask_secret=lambda _: "", existing=existing
    )
    p = data.profile
    assert p["street_address"] == "100 Example St"
    assert p["zip_code"] == "95126"
    assert p["city"] == "San Jose"
    assert p["state"] == "TX"  # not in existing -> resume guess still applies
    assert p["visa_sponsorship"] is True
    assert p["years_of_experience"] == 9
    assert p["gender"] == "Woman"
    assert p["desired_salary"] == "150000"


def test_review_requires_email():
    parsed = _parsed()
    parsed.contact.email = None
    with pytest.raises(SystemExit, match="email is required"):
        interactive_review(parsed, yes=True)


# ---------------------------------------------------------------------------
# save (live PG via db fixture)


def test_save_writes_the_facts_and_is_idempotent(db):
    """Facts only: resume content belongs to a role (`resume add`), so saving the
    facts must neither need nor write a resume."""
    data = interactive_review(_parsed(), yes=True, ask_secret=lambda _: "hunter2")

    uid = save(db, data)
    uid2 = save(db, data)
    assert uid == uid2  # keyed by email

    from common import repo

    p = repo.fetch_profile(db, uid)
    assert p.user["full_name"] == "Jane Doe"
    assert p.profile["city"] == "Austin"
    assert p.profile["first_name"] == "Jane" and p.profile["country"] == "United States"
    assert p.label is None and p.work_experiences == []  # no role added yet
    assert repo.list_base_assets(db, uid) == []
    assert p.login["login_email"] == "jane@x.dev"
    assert p.login["login_password"] == "hunter2"

    # re-running init and skipping the password must keep the stored credential
    data2 = interactive_review(_parsed(), yes=True, ask_secret=lambda _: "")
    save(db, data2)
    assert repo.fetch_profile(db, uid).login["login_password"] == "hunter2"


def test_profile_parser_wiring():
    from cli import main as cli

    args = cli.build_parser().parse_args(["profile", "init", "--email", "a@b.c", "--yes"])
    assert callable(args.func) and args.docx is None  # facts without a resume
    args = cli.build_parser().parse_args(["profile", "init", "--docx", "r.docx"])
    assert args.label == "general"  # the first role, unless named
    args = cli.build_parser().parse_args(["profile", "show"])
    assert callable(args.func)
