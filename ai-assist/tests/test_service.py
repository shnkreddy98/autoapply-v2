"""/answer endpoint orchestration: ownership, context loading, error mapping.
Live PG+MinIO with the agent stubbed out (no LLM)."""

import uuid

import pytest
from fastapi.testclient import TestClient

import qa
import service
from common import keys, repo


def _docx_bytes(tmp_path, lines, table_rows=()):
    """A resume .docx. `table_rows` go into a table, like the real template, which
    keeps employers, titles, dates and the whole education block in tables."""
    from docx import Document

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = text
    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "base.docx"
    doc.save(str(path))
    return path.read_bytes()


def _seed(db, s3, tmp_path, *, with_jd=True, with_resume=True):
    """The `db` fixture TRUNCATEs Postgres between tests but MinIO is never cleaned,
    and job_id is a uuid5 of the URL — so a fixed URL would reuse the previous test's
    JD object. Every seed gets its own URL."""
    with db.cursor() as cur:
        cur.execute(
            "INSERT INTO users (email, default_resume_label) VALUES (%s, 'general') "
            "RETURNING user_id",
            ("q@example.com",),
        )
        user_id = cur.fetchone()["user_id"]
    db.commit()
    if with_resume:
        key = keys.base_resume_key(user_id, "general", "docx", "docx")
        s3.put_bytes(key, _docx_bytes(tmp_path, ["Jane Dev", "Built dbt pipelines"]))
        # variant matters: answering parses the docx, so the pdf base is not a
        # substitute and the lookup pins the format
        repo.insert_resume_asset(db, user_id, "base", key, variant="docx", label="general")
    job_id = repo.add_job_for_user(db, user_id, f"https://example.com/jobs/{uuid.uuid4()}")
    if with_jd:
        s3.put_bytes(keys.jd_key(job_id), b"# Job\nNeed dbt + Airflow.", "text/markdown")
    return user_id, job_id


def _add_labelled_base(db, s3, tmp_path, user_id, label, lines):
    """A role base resume, as `autoapply resume add --label` would create it."""
    key = keys.base_resume_key(user_id, label, "docx", "docx")
    s3.put_bytes(key, _docx_bytes(tmp_path / label, lines))
    return repo.insert_resume_asset(
        db, user_id, "base", key, variant="docx", label=label
    )


@pytest.fixture
def client(monkeypatch):
    service._storage.cache_clear()
    return TestClient(service.app)


def _stub_agent(monkeypatch, captured=None, reply="I built dbt pipelines end to end."):
    async def fake(resume, jd, question, max_words):
        if captured is not None:
            captured.update(resume=resume, jd=jd, question=question, max_words=max_words)
        return reply

    monkeypatch.setattr(qa, "run_agent", fake)


def test_answer_success(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "Why you?"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answer"] == "I built dbt pipelines end to end."
    assert body["word_count"] == 7
    assert body["max_words"] == 100

    assert "dbt pipelines" in seen["resume"]   # resume .docx parsed to text
    assert "Airflow" in seen["jd"]             # captured JD passed through
    assert seen["question"] == "Why you?"
    assert seen["max_words"] == 100


def test_max_words_is_forwarded_and_echoed(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    seen = {}
    _stub_agent(monkeypatch, seen, reply="one two three four five.")

    r = client.post(
        "/answer",
        json={
            "user_id": str(user_id),
            "job_id": str(job_id),
            "question": "Why you?",
            "max_words": 15,
        },
    )
    assert r.status_code == 200
    assert seen["max_words"] == 15
    assert r.json() == {
        "answer": "one two three four five.",
        "word_count": 5,
        "max_words": 15,
        "base_label": "general",
    }


def test_answer_rejects_job_the_user_does_not_own(db, s3, tmp_path, client, monkeypatch):
    _seed(db, s3, tmp_path)
    _stub_agent(monkeypatch)
    with db.cursor() as cur:
        cur.execute(
            "INSERT INTO users (email) VALUES (%s) RETURNING user_id", ("other@e.co",)
        )
        stranger = cur.fetchone()["user_id"]
    db.commit()
    other_job = repo.get_or_create_job(db, f"https://example.com/jobs/{uuid.uuid4()}")

    r = client.post(
        "/answer",
        json={"user_id": str(stranger), "job_id": str(other_job), "question": "Why?"},
    )
    assert r.status_code == 404


def test_answer_without_captured_jd_is_409(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path, with_jd=False)
    _stub_agent(monkeypatch)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "Why?"},
    )
    assert r.status_code == 409
    assert "JD" in r.json()["detail"]


def test_answer_without_base_resume_is_404(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path, with_resume=False)
    _stub_agent(monkeypatch)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "Why?"},
    )
    assert r.status_code == 404
    assert "base resume" in r.json()["detail"]


def test_bearer_token_enforced_when_configured(db, s3, tmp_path, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    _stub_agent(monkeypatch)
    monkeypatch.setattr(service, "TOKEN", "s3cret")
    c = TestClient(service.app)
    payload = {"user_id": str(user_id), "job_id": str(job_id), "question": "Why?"}

    assert c.post("/answer", json=payload).status_code == 401
    assert c.post(
        "/answer", json=payload, headers={"Authorization": "Bearer wrong"}
    ).status_code == 401
    assert c.post(
        "/answer", json=payload, headers={"Authorization": "Bearer s3cret"}
    ).status_code == 200


def test_max_words_out_of_range_is_422(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    _stub_agent(monkeypatch)
    r = client.post(
        "/answer",
        json={
            "user_id": str(user_id),
            "job_id": str(job_id),
            "question": "Why?",
            "max_words": 5,
        },
    )
    assert r.status_code == 422


def test_resume_text_includes_table_content(db, s3, tmp_path, client, monkeypatch):
    """Education and employment live in tables in the real template. Reading only
    `doc.paragraphs` dropped them, and the model answered that the resume has no
    education section — accurately, for the text it was handed."""
    with db.cursor() as cur:
        cur.execute(
            "INSERT INTO users (email, default_resume_label) VALUES (%s, 'general') "
            "RETURNING user_id",
            ("t@example.com",),
        )
        user_id = cur.fetchone()["user_id"]
    db.commit()
    key = keys.base_resume_key(user_id, "general", "docx", "docx")
    s3.put_bytes(
        key,
        _docx_bytes(
            tmp_path,
            ["Jane Dev", "Education"],
            table_rows=[("Master's Degree", "Data Analytics", "Example State University")],
        ),
    )
    repo.insert_resume_asset(db, user_id, "base", key, variant="docx", label="general")
    job_id = repo.add_job_for_user(db, user_id, f"https://example.com/jobs/{uuid.uuid4()}")
    s3.put_bytes(keys.jd_key(job_id), b"# Job\nNeed a degree.", "text/markdown")
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "Degree?"},
    )
    assert r.status_code == 200, r.text
    assert "Example State University" in seen["resume"]
    # order preserved, so the row still reads as sitting under its heading
    assert seen["resume"].index("Education") < seen["resume"].index("Example State")


# ---------------------------------------------------------------------------
# which base resume grounds the answer


def test_answer_follows_the_base_the_job_was_tailored_from(
    db, s3, tmp_path, client, monkeypatch
):
    """The whole point of role bases: an AI/ML posting must be answered from the
    aiml resume. Answering from the default one made the model deny experience
    ("no RAG experience") that is only listed on the role base."""
    user_id, job_id = _seed(db, s3, tmp_path)
    aiml = _add_labelled_base(
        db, s3, tmp_path, user_id, "aiml", ["Jane Dev", "Shipped a RAG pipeline"]
    )
    repo.enqueue_tailoring(db, user_id, job_id, aiml)
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "RAG?"},
    )
    assert r.status_code == 200, r.text
    assert "RAG pipeline" in seen["resume"]
    assert "dbt pipelines" not in seen["resume"]
    assert r.json()["base_label"] == "aiml"


def test_explicit_label_overrides_the_job(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    aiml = _add_labelled_base(
        db, s3, tmp_path, user_id, "aiml", ["Jane Dev", "Shipped a RAG pipeline"]
    )
    repo.enqueue_tailoring(db, user_id, job_id, aiml)
    _add_labelled_base(db, s3, tmp_path, user_id, "de", ["Jane Dev", "Owned Airflow DAGs"])
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={
            "user_id": str(user_id),
            "job_id": str(job_id),
            "question": "Why?",
            "label": "de",
        },
    )
    assert r.status_code == 200, r.text
    assert "Airflow DAGs" in seen["resume"]
    assert r.json()["base_label"] == "de"


def test_label_default_forces_the_default_role(db, s3, tmp_path, client, monkeypatch):
    user_id, job_id = _seed(db, s3, tmp_path)
    aiml = _add_labelled_base(
        db, s3, tmp_path, user_id, "aiml", ["Jane Dev", "Shipped a RAG pipeline"]
    )
    repo.enqueue_tailoring(db, user_id, job_id, aiml)
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={
            "user_id": str(user_id),
            "job_id": str(job_id),
            "question": "Why?",
            "label": "default",
        },
    )
    assert r.status_code == 200, r.text
    assert "dbt pipelines" in seen["resume"]
    assert r.json()["base_label"] == "general"


def test_unknown_label_falls_back_to_the_default_base(
    db, s3, tmp_path, client, monkeypatch
):
    """A label with nothing uploaded still answers — from the default role — and
    says which role that was, rather than 404ing mid-application."""
    user_id, job_id = _seed(db, s3, tmp_path)
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={
            "user_id": str(user_id),
            "job_id": str(job_id),
            "question": "Why?",
            "label": "nosuch",
        },
    )
    assert r.status_code == 200, r.text
    assert "dbt pipelines" in seen["resume"]
    assert r.json()["base_label"] == "general"


def test_untailored_job_uses_the_default_base(db, s3, tmp_path, client, monkeypatch):
    """No tailoring row -> nothing to follow -> the user's default role."""
    user_id, job_id = _seed(db, s3, tmp_path)
    _add_labelled_base(db, s3, tmp_path, user_id, "aiml", ["Jane Dev", "RAG pipeline"])
    seen = {}
    _stub_agent(monkeypatch, seen)

    r = client.post(
        "/answer",
        json={"user_id": str(user_id), "job_id": str(job_id), "question": "Why?"},
    )
    assert r.status_code == 200, r.text
    assert "dbt pipelines" in seen["resume"]
    assert r.json()["base_label"] == "general"
