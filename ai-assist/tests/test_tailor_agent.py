"""ResumeTailorAgent: the prompt, the edit budget (summary + TAILOR_BULLETS bullets +
TAILOR_SKILL_LINES Skills lines, counted only when a replacement actually lands) and
the replace tool on a real .docx. No LLM, no DB."""

import pytest

from agents import ResumeTailorAgent
from prompts import system_prompt_tailor


class FakeDoc:
    """DocumentTools stand-in: replace() answers from a script, records calls."""

    def __init__(self, results):
        self.results = list(results)
        self.calls = []
        self.document = None

    async def replace(self, args):
        self.calls.append(args)
        return self.results.pop(0)


def test_prompt_budget_counts_summary_bullets_and_skills():
    p = system_prompt_tailor(5, 2)
    assert "make 8 targeted replacements" in p
    assert "5 bullet point rewrite(s)" in p and "up to 2 Skills line update(s)" in p
    assert "{EDITS}" not in p and "{BULLETS}" not in p and "{SKILLS}" not in p
    assert "make 2 targeted replacements" in system_prompt_tailor(1, 0)


def test_prompt_shapes_the_summary_and_lets_jd_tools_into_skills():
    p = system_prompt_tailor(5, 2)
    # the summary opens the same way for every JD, whatever its title
    assert '"Software Engineer with experience in <' in p
    # tools missing from the one-page resume are not treated as missing skills
    assert "a tool missing from the resume is NOT evidence the candidate lacks it" in p
    # Skills used to be frozen; added tools must be backed by a bullet
    assert "Never modify the Skills section" not in p
    assert "must appear in at least one rewritten bullet" in p
    assert "Never paste JD sentences or phrases" in p
    # page length was dropped as a constraint (2026-10-29): additions may grow the resume
    assert "1 page" not in p and "same length" not in p


@pytest.mark.asyncio
async def test_only_landed_replacements_count_and_extras_are_refused():
    ok = "Successfully replaced"
    miss = "ERROR: search_text not found in document. Make sure ..."
    doc = FakeDoc([ok, miss, ok, ok])
    agent = ResumeTailorAgent(doc, bullets=1, skills=1)  # 3 edits: summary + bullet + skills
    replace = agent.tool_functions["replace"]

    for _ in range(3):
        await replace("args")
    assert agent._tool_choice() == "auto"  # the miss did not use up an edit
    await replace("args")
    assert agent._tool_choice() == "none"

    # a fourth call in the same response never reaches the document
    out = await replace("args")
    assert out.startswith("Refused")
    assert len(doc.calls) == 4


def _skills_doc(path):
    import docx

    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Languages:").bold = True
    p.add_run(" Python, SQL, Bash")
    d.add_paragraph("Built Airflow DAGs moving 2 TB/day.")
    d.save(str(path))
    return str(path)


@pytest.mark.asyncio
async def test_replace_inside_one_run_keeps_the_bold_label(tmp_path):
    """A Skills line is a bold "Languages:" run plus a plain tool list. Replacing in
    the list must not restyle the label run or make the list bold."""
    import docx

    from tools import DocumentTools, ReplaceArgs

    path = _skills_doc(tmp_path / "r.docx")
    out = await DocumentTools(path).replace(
        ReplaceArgs(search_text="Python, SQL, Bash", replace_text="Python, Go, SQL, Bash")
    )
    assert out == "Successfully replaced"
    runs = docx.Document(path).paragraphs[0].runs
    assert [(r.text, bool(r.bold)) for r in runs] == [
        ("Languages:", True), (" Python, Go, SQL, Bash", False),
    ]


@pytest.mark.asyncio
async def test_replace_spanning_runs_still_works(tmp_path):
    import docx

    from tools import DocumentTools, ReplaceArgs

    path = _skills_doc(tmp_path / "r.docx")
    out = await DocumentTools(path).replace(
        ReplaceArgs(search_text="Languages: Python", replace_text="Languages: Go")
    )
    assert out == "Successfully replaced"
    assert docx.Document(path).paragraphs[0].text == "Languages: Go, SQL, Bash"
