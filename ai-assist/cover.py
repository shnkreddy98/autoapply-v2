"""Cover letters for a user x job: written by CoverLetterAgent, laid out as a .docx,
rendered to PDF, and stored next to the job (common.keys.cover_letter_key).

Grounded in the same base resume `ask` uses (qa.choose_base_asset: the role the job
was tailored from, unless a label is forced), so a letter for an AI/ML posting
draws on the aiml resume rather than the default one.

The model writes only the body. The contact header, date, salutation and sign-off
come from the users row, so a letter can never carry a made-up phone number or a
misspelt name.
"""

from __future__ import annotations

import shutil
import tempfile
from datetime import date
from pathlib import Path

from common import keys

import qa

MAX_WORDS_DEFAULT = 350


def letter_header(conn, user_id: str, job_id: str) -> dict:
    """Name and contact from the users row; company and title from the job."""
    with conn.cursor() as cur:
        cur.execute(
            "SELECT full_name, email, phone, linkedin, location FROM users WHERE user_id = %s",
            (user_id,),
        )
        user = cur.fetchone() or {}
        cur.execute("SELECT company, title FROM jobs WHERE job_id = %s", (job_id,))
        job = cur.fetchone() or {}
    return {**user, "company": job.get("company"), "title": job.get("title")}


def letter_text(paragraphs: list[str], header: dict, today: date | None = None) -> str:
    """The whole letter as plain text: what the PDF shows, and what gets pasted into
    a "cover letter" text box."""
    name = header.get("full_name") or ""
    contact = " | ".join(
        str(header[k]) for k in ("location", "phone", "email", "linkedin") if header.get(k)
    )
    salutation = (
        f"Dear {header['company']} Hiring Team," if header.get("company") else "Dear Hiring Manager,"
    )
    parts = [name, contact, (today or date.today()).strftime("%B %d, %Y"), salutation]
    parts += paragraphs
    parts += ["Sincerely,", name]
    return "\n\n".join(p for p in parts if p)


def build_docx(paragraphs: list[str], header: dict, path: str, today: date | None = None) -> str:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    name = doc.add_paragraph()
    name.add_run(header.get("full_name") or "").bold = True
    name.runs[0].font.size = Pt(14)
    # the rest of the letter is exactly letter_text minus the name line, so the PDF
    # and the pasted text never drift apart
    for block in letter_text(paragraphs, header, today).split("\n\n")[1:]:
        doc.add_paragraph(block)
    doc.save(path)
    return path


def load(conn, storage, user_id: str, job_id: str, label: str | None = None) -> dict:
    """Everything the letter needs from Postgres/MinIO, after qa's ownership check.
    Separate from `write_and_store` so the caller can release the connection before
    the LLM call."""
    resume, jd, used = qa.load_context(conn, storage, user_id, job_id, label)
    return {
        "user_id": user_id,
        "job_id": job_id,
        "resume": resume,
        "jd": jd,
        "base_label": used,
        "header": letter_header(conn, user_id, job_id),
    }


async def write_and_store(storage, ctx: dict, max_words: int = MAX_WORDS_DEFAULT) -> dict:
    """Write the letter from `load`'s context, store pdf + txt, return {"text",
    "word_count", "base_label", "pdf_key", "txt_key"}. Raises RuntimeError."""
    from word import convert_docx_to_pdf

    paragraphs = await run_agent(ctx["resume"], ctx["jd"], max_words)
    if not paragraphs:
        raise RuntimeError("cover letter came back empty")

    work = Path(tempfile.mkdtemp(prefix="cover-"))
    try:
        docx_path = build_docx(paragraphs, ctx["header"], str(work / "cover_letter.docx"))
        pdf = await convert_docx_to_pdf(docx_path)
        if not pdf:
            raise RuntimeError("cover letter docx->pdf failed (is LibreOffice installed?)")
        text = letter_text(paragraphs, ctx["header"])
        pdf_key = keys.cover_letter_key(ctx["user_id"], ctx["job_id"], "pdf")
        txt_key = keys.cover_letter_key(ctx["user_id"], ctx["job_id"], "txt")
        storage.put_file(pdf_key, pdf)
        storage.put_bytes(txt_key, text.encode("utf-8"), "text/plain; charset=utf-8")
    finally:
        shutil.rmtree(work, ignore_errors=True)

    return {
        "text": text,
        "word_count": sum(len(p.split()) for p in paragraphs),
        "base_label": ctx["base_label"],
        "pdf_key": pdf_key,
        "txt_key": txt_key,
    }


async def run_agent(resume: str, jd: str, max_words: int) -> list[str]:
    """Indirection so tests can swap the LLM out. Imported lazily: `agents` pulls in
    `env`, which refuses to load without an API key."""
    from agents import CoverLetterAgent

    return await CoverLetterAgent().write(resume, jd, max_words=max_words)
